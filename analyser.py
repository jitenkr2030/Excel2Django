"""Analyze Excel files."""
from pathlib import Path
from openpyxl import load_workbook
from helpers import slug, cname

def analyse(source):
    src = Path(source)
    if src.is_file(): files = [src]
    elif src.is_dir():
        files = sorted(p for p in src.iterdir() if p.suffix.lower() in (".xlsx",".xlsm") and not p.name.startswith("~$"))
    else: raise FileNotFoundError(f"Not found: {src}")
    if not files: raise ValueError("No Excel files")
    result, used = [], {"User"}
    for f in files:
        wb = load_workbook(f, read_only=True, data_only=True)
        for ws in wb.worksheets:
            hdrs = next(ws.iter_rows(min_row=1,max_row=1,values_only=True), None)
            if not hdrs: continue
            sample = list(ws.iter_rows(min_row=2, max_row=min(ws.max_row or 2, 31), values_only=True))
            fields, used_f = [], set()
            for i, raw in enumerate(hdrs):
                lab = str(raw or f"Col{i+1}").strip()
                fld = slug(lab); base, s = fld, 2
                while fld in used_f or fld == "id": fld, s = f"{base}_{s}", s+1
                used_f.add(fld)
                vals = [r[i] for r in sample if i < len(r) and r[i] is not None]
                if vals and all(isinstance(v,bool) for v in vals): t = "BooleanField"
                elif vals and all(isinstance(v,int) and not isinstance(v,bool) for v in vals): t = "BigIntegerField"
                elif vals and all(isinstance(v,(int,float)) and not isinstance(v,bool) for v in vals): t = "FloatField"
                elif vals and all(hasattr(v,"year") and hasattr(v,"month") and hasattr(v,"day") for v in vals): t = "DateField"
                else: t = "TextField"
                fields.append({"name":fld,"label":lab,"type":t})
            bm = cname(ws.title); m, s = bm, 2
            while m in used: m, s = f"{bm}{s}", s+1
            used.add(m)
            result.append({"file":f.name,"sheet":ws.title,"model":m,"fields":fields,"rows":max((ws.max_row or 1)-1,0)})
        wb.close()
    if not result: raise ValueError("No worksheets found")
    return result
