"""Core engine to generate Excel files from schemas."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from gen_excel_schemas import SCHEMAS as S1
from gen_excel_schemas2 import SCHEMAS as S2
from gen_excel_schemas3 import SCHEMAS as S3
from gen_excel_cat1 import SCHEMAS as C1
from gen_excel_cat2 import SCHEMAS as C2
from gen_excel_cat3 import SCHEMAS as C3
from gen_excel_biztypes import BIZ_TYPES, BIZ_CATEGORIES, get_schema_for_business, list_businesses

ALL_SCHEMAS = {**S1, **S2, **S3, **C1, **C2, **C3}

HDR_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HDR_FILL = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")

def _generate(domain, rows, output, title_override=None):
    if domain not in ALL_SCHEMAS:
        return False
    schema = ALL_SCHEMAS[domain]
    title = title_override or schema["title"]
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, config in schema["sheets"].items():
        ws = wb.create_sheet(title=sheet_name[:31])
        cols = config["cols"]
        gen_func = config["gen"]
        for col_idx, col_name in enumerate(cols, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(col_idx)].width = max(12, len(col_name)+4)
        for row_idx in range(1, rows+1):
            data = gen_func(row_idx)
            for col_idx, val in enumerate(data, 1):
                ws.cell(row=row_idx+1, column=col_idx, value=val)
    out_path = Path(output)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    total_sheets = len(schema["sheets"])
    print(f"Generated: {out_path}")
    print(f"  Title: {title} | Domain: {domain} | Sheets: {total_sheets} | Rows/sheet: {rows}")
    for sn, sc in schema["sheets"].items():
        print(f"    {sn}: {len(sc['cols'])} cols, {rows} rows")
    return True

def generate_excel(domain, rows, output):
    return _generate(domain, rows, output)

def generate_business(biz_name, rows, output):
    cat, title = get_schema_for_business(biz_name)
    if not cat:
        print(f"Unknown business: {biz_name}")
        return False
    return _generate(cat, rows, output, title_override=title)

def generate_all(rows, output_dir="generated_excel"):
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    for domain, schema in ALL_SCHEMAS.items():
        fname = f"{schema['title']}_Demo.xlsx"
        _generate(domain, rows, out / fname)
        print()
    print(f"All {len(ALL_SCHEMAS)} domains generated in {out}/")

def generate_all_businesses(rows, output_dir="generated_businesses"):
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    for biz_name, cat in BIZ_TYPES.items():
        _, title = get_schema_for_business(biz_name)
        fname = f"{title}.xlsx"
        _generate(cat, rows, out / fname, title_override=title)
    print(f"All {len(BIZ_TYPES)} businesses generated in {out}/")

def list_domains():
    print(f"Available domains ({len(ALL_SCHEMAS)}):")
    for domain, schema in ALL_SCHEMAS.items():
        sheets = schema["sheets"]
        cols = sum(len(s["cols"]) for s in sheets.values())
        print(f"  {domain:15s} | {schema['title']:45s} | {len(sheets):2d} sheets | {cols:3d} cols")

def list_all():
    list_domains()
    print()
    list_businesses()
